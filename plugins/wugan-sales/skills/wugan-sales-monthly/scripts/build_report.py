"""五感月度销售表：raw json（fetch_month.js 输出）+ rules.yaml → 追加到工作簿的当月 sheet。

退出码：0 成功；2 有未登记的规格选项（需改 rules.yaml）；3 对数不平；4 sheet 已存在；5 有未确认的大额消费。
"""
import argparse
import collections
import datetime as dt
import json
import re
import pathlib
import shutil
import sys

import openpyxl
import openpyxl.utils
import yaml
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADER = ["分类", "商品名称", "规格", "单价", "实际销量", "实际销售额\n(不含退款、优惠)", "实际销售额占比"]  # F 列手动换行，避免断在「优惠」中间
BLANK_SPEC = "(空白)"


def cents(x):
    return round(float(x or 0) * 100)


def load_raw(path):
    d = json.loads(pathlib.Path(path).read_text(encoding="utf-8"))
    if isinstance(d, str):  # Playwright 落盘有时是 JSON 字符串套 JSON
        d = json.loads(d)
    if "error" in d:
        sys.exit(f"raw 文件是取数失败的结果：{d}")
    return d


def clean(raw, rules):
    spec_rules = rules.get("spec_rules") or {}
    unknown = collections.Counter()
    out = []
    for r in raw["list"]:
        title = (r["title"] or "").strip()
        cat = (rules.get("reclassify") or {}).get(title, r["classify_title"])
        orig_cat = r["classify_title"]
        title = (rules.get("product_rename") or {}).get(title, title)
        spec = None
        sr = spec_rules.get(orig_cat) or spec_rules.get(cat)
        if sr:
            opts = [o.strip() for o in r["key_list"] if o and o.strip()]
            for o in opts:
                if o not in sr.get("keep", []) and o not in sr.get("drop", []):
                    unknown[(orig_cat, o)] += 1
            kept = [o for o in opts if o in sr.get("keep", [])]
            if kept:
                spec = (rules.get("spec_rename") or {}).get(kept[0], kept[0])
        out.append({
            "cat": cat, "orig_cat": orig_cat, "orig_title": r["title"], "title": title, "spec": spec,
            "one_price": float(r["one_price"] or 0), "real_num": int(float(r["real_num"] or 0)),
            "real_total": float(r["real_total"] or 0), "discount": float(r["discount_amount"] or 0),
            "sum_total": float(r["sum_total"] or 0), "options": "|".join(r["key_list"]),
        })
    return out, unknown


def aggregate(rows, cat_order):
    """三层：分类 → 商品 → 规格。单价 = 明细行 one_price 的算术平均。"""
    def node():
        return {"prices": [], "qty": 0, "amt_c": 0, "kids": collections.OrderedDict()}

    root = collections.OrderedDict()
    for r in rows:
        c = root.setdefault(r["cat"], node())
        p = c["kids"].setdefault(r["title"], node())
        targets = [c, p]
        if r["spec"] is not None or any(x["spec"] for x in rows if x["cat"] == r["cat"] and x["title"] == r["title"]):
            targets.append(p["kids"].setdefault(r["spec"] or BLANK_SPEC, node()))
        for n in targets:
            n["prices"].append(r["one_price"])
            n["qty"] += r["real_num"]
            n["amt_c"] += cents(r["real_total"])

    def order_key(name):
        return (cat_order.index(name) if name in cat_order else len(cat_order), name)

    def sort_kids(n):
        items = sorted(n["kids"].items(), key=lambda kv: (kv[0] == BLANK_SPEC, -kv[1]["amt_c"], kv[0]))
        n["kids"] = collections.OrderedDict(items)
        for _, k in items:
            sort_kids(k)

    ordered = collections.OrderedDict(sorted(root.items(), key=lambda kv: order_key(kv[0])))
    for n in ordered.values():
        sort_kids(n)
    return ordered


def flatten(tree):
    total_c = sum(n["amt_c"] for n in tree.values())
    lines = []
    for cname, c in tree.items():
        lines.append((0, [cname, None, None], c))
        for pname, p in c["kids"].items():
            lines.append((1, [None, pname, None], p))
            for sname, s in p["kids"].items():
                lines.append((2, [None, None, sname], s))
    out = []
    for level, names, n in lines:
        price = sum(n["prices"]) / len(n["prices"]) if n["prices"] else None
        share = n["amt_c"] / total_c if total_c else 0
        out.append((level, names + [price, n["qty"], n["amt_c"] / 100, share]))
    all_prices = [x for n in tree.values() for x in n["prices"]]
    out.append((-1, ["总计", None, None, sum(all_prices) / len(all_prices) if all_prices else None,
                     sum(n["qty"] for n in tree.values()), total_c / 100, 1 if total_c else 0]))
    return out


FILE_PATTERN = "{year}年五感销售数据统计.xlsx"
SHEET_RE = re.compile(r"^(\d{4})年(\d{1,2})月(-.*)?$")


def sheet_month(title):
    """「2026年9月」「2026年9月-异常与调整」→ (2026, 9)；旧版无年份 sheet（「8月」）返回 None。"""
    mt = SHEET_RE.match(title)
    return (int(mt.group(1)), int(mt.group(2))) if mt else None


def history_names(books, target):
    """target 月之前所有 sheet 里出现过的 商品名 / (商品名, 规格)，用来识别新品。
    books：同目录下所有年度文件（含旧版无年份文件）；旧版 sheet 无年份，视为更早。"""
    prods, specs = set(), set()
    for ws in (w for wb in books for w in wb.worksheets):
        sm = sheet_month(ws.title)
        if sm is not None and sm >= target:
            continue
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        rows = [r for r in rows if any(v is not None for v in r)]
        if not rows:
            continue
        head = [str(v).strip() if v is not None else "" for v in rows[0]]
        if "商品名称" not in head or "规格" not in head:
            continue
        bi, ci = head.index("商品名称"), head.index("规格")
        cur = None
        for r in rows[1:]:
            if r[bi]:
                cur = str(r[bi]).strip()
                prods.add(cur)
            elif r[ci] and cur:
                specs.add((cur, str(r[ci]).strip()))
    return prods, specs


def sort_month_sheets(wb):
    """按月份排列：主表 → -1元计价 → -异常与调整；无年份的旧 sheet 留在前面。"""
    suffix_rank = {None: 0, "-1元计价": 1, "-异常与调整": 2}

    def key(ws):
        mt = SHEET_RE.match(ws.title)
        if not mt:
            return (0, 0, 0, 0)
        return (1, int(mt.group(1)), int(mt.group(2)), suffix_rank.get(mt.group(3), 3))
    wb._sheets.sort(key=key)


# ---- 样式 ----
FONT = "微软雅黑"
C_HEAD, C_CAT, C_TOTAL, C_LINE, C_MUTED, C_BAR = "5B4636", "F3EDE6", "E8DED3", "E3DCD4", "8C8C8C", "C9A27E"
THIN = Side(style="thin", color=C_LINE)
ISSUE_FILL = {"新": "DDEBF7", "1元": "FCE4D6", "0元": "EDEDED", "改": "E2EFDA", "去首尾": "E2EFDA"}


def style_header(row):
    for c in row:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=C_HEAD)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_note(ws, text, ncols):
    ws.append([])
    ws.append([text])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    c = ws.cell(r, 1)
    c.font = Font(name=FONT, italic=True, size=9, color=C_MUTED)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 32


def write_sheet(ws, lines, note=None):
    ws.sheet_view.showGridLines = False
    ws.append(HEADER)
    style_header(ws[1])
    ws.row_dimensions[1].height = 40
    cat_rows = []
    for level, vals in lines:
        ws.append(vals)
        r = ws.max_row
        row = ws[r]
        if level == 0:
            font, fill = Font(name=FONT, bold=True, size=11), PatternFill("solid", fgColor=C_CAT)
            cat_rows.append(r)
        elif level == -1:
            font, fill = Font(name=FONT, bold=True, size=11), PatternFill("solid", fgColor=C_TOTAL)
        elif level == 2:
            font, fill = Font(name=FONT, size=9, color=C_MUTED), None
        else:
            font, fill = Font(name=FONT, size=10), None
        for c in row:
            c.font = font
            if fill:
                c.fill = fill
            c.border = Border(bottom=THIN, top=Side(style="medium", color=C_HEAD) if level == -1 else None)
            c.alignment = Alignment(vertical="center")
        if level == 2:
            row[2].alignment = Alignment(indent=1, vertical="center")
        if level in (1, 2):  # 分级折叠：1=只看分类，2=到商品，3=全部
            ws.row_dimensions[r].outline_level = level
        ws.row_dimensions[r].height = 22 if level in (0, -1) else 18
        row[3].number_format = "0.00"
        row[4].number_format = "#,##0"
        row[5].number_format = "#,##0.00"
        row[6].number_format = "0.0%"
    ws.sheet_properties.outlinePr.summaryBelow = False
    if cat_rows:  # 分类占比加数据条（分类之间可比，合计 100%）
        ws.conditional_formatting.add(" ".join(f"G{r}" for r in cat_rows),
                                      DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=C_BAR))
    if note:
        add_note(ws, note, len(HEADER))
    for col, w in zip("ABCDEFG", [22, 34, 16, 10, 12, 22, 19]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def write_issues(ws, issues, recon):
    ws.sheet_view.showGridLines = False
    ws.append(["类型", "分类", "商品", "规格/选项", "说明", "实际销量", "实际销售额"])
    style_header(ws[1])
    ws.row_dimensions[1].height = 28
    for i in issues:
        ws.append(i)
        row = ws[ws.max_row]
        tag = next((v for k, v in ISSUE_FILL.items() if str(i[0]).startswith(k)), None)
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.border = Border(bottom=THIN)
            c.alignment = Alignment(vertical="center", wrap_text=True)
        if tag:
            row[0].fill = PatternFill("solid", fgColor=tag)
        row[6].number_format = "#,##0.00"
    ws.append([])
    ws.append(["对数结果"])
    ws.cell(ws.max_row, 1).font = Font(name=FONT, bold=True, size=11, color=C_HEAD)
    for label, value in recon:
        ws.append([label, value])
        ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=5)
        for c in ws[ws.max_row]:
            c.font = Font(name=FONT, size=10)
            c.border = Border(bottom=THIN)
    for col, w in zip("ABCDEFG", [24, 18, 32, 26, 34, 12, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


SUMMARY_SHEET = "年度汇总"
OPS_SHEET = "月度经营数据"
OPS_COLS = ["月份", "营业天数", "订单数", "新客单数", "老客单数", "订单实收", "余额支付", "储值充值", "充值笔数",
            "余额消费", "期末储值余额", "主表销售额", "主表件数", "剔除异常"]


def _f(x):
    try:
        return float(x or 0)
    except (TypeError, ValueError):
        return 0.0


def apply_anomalies(raw, rules):
    """单笔 ≥ 阈值的已支付订单 / 余额商家扣除：已确认剔除的扣掉，已确认保留的放过，其余列为待确认。
    返回 (调整后的商品明细, 剔除汇总, 待确认清单, 异常 sheet 行)。"""
    an = rules.get("anomaly") or {}
    thr = float(an.get("threshold", 10000))
    month = raw["month"]
    ex_orders = {(o["pay_time"], round(float(o["total"]))): o for o in an.get("excluded_orders") or []}
    pairs = an.get("excluded_balance_pairs") or []
    ex_deds = {(pr["deduction"]["time"], round(float(pr["deduction"]["amount"]))): pr for pr in pairs}
    kept = {(k.get("pay_time") or k.get("time"), round(float(k.get("total") or k.get("amount")))) for k in an.get("kept") or []}

    items = json.loads(json.dumps(raw["list"]))
    ex = {"order_n": 0, "order_amt": 0.0, "ded_amt": 0.0, "topup_amt": 0.0, "topup_n": 0, "notes": []}
    pending, issue_rows = [], []

    for o in raw.get("big_orders") or []:
        key = (o["pay_time"], round(_f(o["total"])))
        if key in ex_orders:
            ex["order_n"] += 1
            ex["order_amt"] += _f(o["total"])
            ex["notes"].append(f"订单 {o['pay_time']} {_f(o['total']):,.0f}")
            issue_rows.append(["剔除异常订单", "", "、".join(g["title"] for g in o["goods"]), o["pay_time"],
                               ex_orders[key].get("note", ""), "", -_f(o["total"])])
            for g in o["goods"]:  # 从商品明细里扣掉这笔订单的商品
                cands = [r for r in items if (r["title"] or "").strip() == (g["title"] or "").strip()]
                if not cands:
                    sys.exit(f"异常订单 {o['pay_time']} 的商品「{g['title']}」在销售汇总里找不到，无法剔除")
                r = max(cands, key=lambda r: _f(r["real_total"]))
                for k, v in (("real_num", _f(g["num"])), ("sum_num", _f(g["num"])),
                             ("real_total", _f(g["total"])), ("sum_total", _f(g["total"]))):
                    r[k] = round(_f(r[k]) - v, 2)
        elif key not in kept:
            pending.append(f"订单  {o['pay_time']}  {_f(o['total']):,.2f} 元  {o.get('pay_type','')}  "
                           f"{'后台下单' if o.get('admin_order') else '顾客下单'}  商品：{'、'.join(g['title'] + '×' + str(g['num']) for g in o['goods'])}")

    for l in raw.get("ledger") or []:
        amt = _f(l["amount"])
        if not str(l["remark"]).startswith("商家扣除") or abs(amt) < thr:
            continue
        key = (l["time"], round(-abs(amt)))
        if key in ex_deds:
            ex["ded_amt"] += abs(amt)
            ex["notes"].append(f"商家扣除 {l['time']} {abs(amt):,.0f}")
            issue_rows.append(["剔除异常扣除(成对)", "", l["remark"], l["time"], ex_deds[key].get("note", ""), "", -abs(amt)])
        elif key not in kept:
            pending.append(f"余额扣除  {l['time']}  {amt:,.2f} 元  {l['remark']}  —— 请连同它对应的充值一起确认")
    for pr in pairs:  # 成对剔除的充值按充值发生的月份扣
        for t in pr["topups"]:
            if t["time"].startswith(month):
                paid = _f(t.get("paid", t["amount"]))  # 营业报表「储值充值」是实付口径，不含赠送
                ex["topup_amt"] += paid
                ex["topup_n"] += 1 if paid else 0
                issue_rows.append(["剔除异常充值(成对)", "", "余额充值", t["time"],
                                   pr.get("note", "") + (f"（到账 {_f(t['amount']):,.0f}，实付 {paid:,.0f}）" if paid != _f(t["amount"]) else ""),
                                   "", -paid])
    return items, ex, pending, issue_rows


def ops_row(raw, ex, main_qty, main_amt):
    """营业报表 + 余额流水 → 当月经营数据（已剔除异常）。旧版 raw 没有这些字段时返回 None。"""
    days = raw.get("report_days")
    if not days:
        return None
    t = collections.Counter()
    for d in days:
        for k, v in d.items():
            if k != "date":
                t[k] += _f(v)
    paid_balance = t["order_total_balance"] - t["order_refund_balance"]
    summ = raw.get("ledger_summary") or {}
    return [raw["month"], len(days), t["order_count"] - ex["order_n"], t["new_user_order_count"], t["patrons_user_order_count"],
            round(t["order_total"] - t["order_refund_total"] - ex["order_amt"], 2), round(paid_balance, 2),
            round(t["topup_total"] - ex["topup_amt"], 2), t["topup_count"] - ex["topup_n"],
            round(paid_balance + (t["cashier_total_balance"] if "cashier_total_balance" in t else t["cashier_total"]) - ex["ded_amt"], 2),
            _f(summ.get("amount_after")) if summ else None, round(main_amt, 2), main_qty, "；".join(ex["notes"])]


OPS_COUNT_COLS = {"营业天数", "订单数", "新客单数", "老客单数", "充值笔数", "主表件数"}
OPS_MONEY_COLS = {"订单实收", "余额支付", "储值充值", "余额消费", "期末储值余额", "主表销售额"}
MONTH_KEY_RE = re.compile(r"^\d{4}-\d{2}$")


def text_width(text, cjk=2.2, other=1.15):
    """估算单行显示所需列宽：中文约 2.2、其他字符约 1.15 个宽度单位。"""
    return sum(cjk if ord(ch) > 0x2E80 else other for ch in str(text))


def fit_widths(ws, headers, data_rows, pad=3, min_w=8, max_w=60):
    """表头单行不换行：列宽 ≥ 表头文字宽度，也照顾数据宽度。"""
    for i, h in enumerate(headers):
        need = text_width(h) + pad
        for r in data_rows:
            v = r[i] if i < len(r) else None
            if v is None:
                continue
            shown = f"{v:,.2f}" if isinstance(v, float) else (f"{v:,}" if isinstance(v, int) else str(v))
            need = max(need, text_width(shown) + 2)
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = max(min_w, min(max_w, need))


def write_ops_sheet(wb, rows):
    """rows：{"YYYY-MM": [与 OPS_COLS 对应的值]}。整张表按固定格式重写，不改数字。"""
    pos = wb.sheetnames.index(OPS_SHEET) if OPS_SHEET in wb.sheetnames else None
    if pos is not None:
        del wb[OPS_SHEET]
    ws = wb.create_sheet(OPS_SHEET, pos)
    ws.sheet_view.showGridLines = False
    ws.append(OPS_COLS)
    style_header(ws[1])
    for c in ws[1]:
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    ws.row_dimensions[1].height = 26
    data = [rows[k] for k in sorted(rows)]
    for vals in data:
        ws.append(vals)
        for col, c in zip(OPS_COLS, ws[ws.max_row]):
            c.font = Font(name=FONT, size=10)
            c.border = Border(bottom=THIN)
            c.alignment = Alignment(vertical="center")
            if col in OPS_COUNT_COLS:
                c.number_format = "#,##0"
            elif col in OPS_MONEY_COLS:
                c.number_format = "#,##0.00"
        ws.row_dimensions[ws.max_row].height = 20
    fit_widths(ws, OPS_COLS, data)
    ws.freeze_panes = "B2"
    add_note(ws, "口径：订单数/订单实收来自营业报表（订单额 − 退款），已剔除 rules.yaml 里确认的异常订单；"
                 "余额消费 = 余额支付 + 商家扣除（剔除成对异常）；储值充值为实付金额（不含赠送，剔除成对异常）；"
                 "期末储值余额取余额流水月末值；主表销售额来自销售汇总。", len(OPS_COLS))


def read_ops_rows(wb):
    """只读第一列是 YYYY-MM 的数据行（跳过表头和口径说明）。"""
    if OPS_SHEET not in wb.sheetnames:
        return {}
    return {str(r[0]): list(r[:len(OPS_COLS)]) for r in wb[OPS_SHEET].iter_rows(min_row=2, values_only=True)
            if r and r[0] is not None and MONTH_KEY_RE.match(str(r[0]))}


def upsert_ops(wb, row):
    rows = read_ops_rows(wb)
    rows[row[0]] = row
    write_ops_sheet(wb, rows)


def read_ops(wb):
    if OPS_SHEET not in wb.sheetnames:
        return {}
    out = {}
    for k, r in read_ops_rows(wb).items():
        out[k] = dict(zip(OPS_COLS, r))
    return out


def read_month_sheet(ws):
    """从月度主表读出 {分类: (销量, 销售额)} 和 总计 (销量, 销售额)。"""
    cats, total = collections.OrderedDict(), (0, 0.0)
    for r in ws.iter_rows(min_row=2, values_only=True):
        name = r[0]
        if not name or str(name).startswith("口径") or str(name).startswith("补差价"):
            continue
        if name == "总计":
            total = (r[4] or 0, r[5] or 0.0)
        else:
            cats[name] = (r[4] or 0, r[5] or 0.0)
    return cats, total


def write_year_summary(wb, year, cat_order):
    """第一页：分类 × 月份 的销售额 / 销量矩阵。只读各月主表，可随时重建。"""
    months = sorted(sheet_month(t)[1] for t in wb.sheetnames
                    if sheet_month(t) and sheet_month(t)[0] == year and not SHEET_RE.match(t).group(3))
    if not months:
        return
    data, totals, one_yuan = {}, {}, {}
    for mo in months:
        title = f"{year}年{mo}月"
        data[mo], totals[mo] = read_month_sheet(wb[title])
        oy = f"{title}-1元计价"
        one_yuan[mo] = read_month_sheet(wb[oy])[1] if oy in wb.sheetnames else (0, 0.0)
    seen = {c for mo in months for c in data[mo]}
    cats = [c for c in cat_order if c in seen] + sorted(seen - set(cat_order))

    if SUMMARY_SHEET in wb.sheetnames:
        del wb[SUMMARY_SHEET]
    ws = wb.create_sheet(SUMMARY_SHEET, 0)
    ws.sheet_view.showGridLines = False
    ncol = 1 + len(months) + 2
    ws.append([f"{year}年 五感销售年度汇总"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=C_HEAD)
    ws.row_dimensions[1].height = 28
    ws.append([f"数据来自各月主表（支付时间口径，单位：元/件）；主表不含 1 元计价商品；更新于 {dt.datetime.now():%Y-%m-%d %H:%M}"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color=C_MUTED)

    def section(title, idx, fmt, with_share):
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(name=FONT, bold=True, size=12, color=C_HEAD)
        ws.append(["分类"] + [f"{mo}月" for mo in months] + ["全年合计", "全年占比" if with_share else ""])
        style_header(ws[ws.max_row][:ncol - (0 if with_share else 1)])
        ws.row_dimensions[ws.max_row].height = 26
        grand = sum(totals[mo][idx] for mo in months)
        first = ws.max_row + 1

        def put(label, vals, kind):
            yr = sum(vals)
            share = (yr / grand if grand else 0) if with_share and kind == "cat" else None
            ws.append([label] + vals + [yr, share])
            row = ws[ws.max_row]
            bold = kind in ("total", "raw")
            fill = PatternFill("solid", fgColor=C_TOTAL) if kind == "total" else None
            for c in row[:ncol]:
                c.font = Font(name=FONT, bold=bold, size=10, color=C_MUTED if kind == "aux" else None)
                c.border = Border(bottom=THIN, top=Side(style="medium", color=C_HEAD) if kind == "total" else None)
                c.alignment = Alignment(vertical="center")
                if fill:
                    c.fill = fill
            for c in row[1:ncol - 1]:
                c.number_format = fmt
            row[ncol - 2].font = Font(name=FONT, bold=True, size=10)
            row[ncol - 1].number_format = "0.0%"
            ws.row_dimensions[ws.max_row].height = 20

        for c in cats:
            put(c, [data[mo].get(c, (0, 0.0))[idx] for mo in months], "cat")
        last_cat = ws.max_row
        put("主表合计", [totals[mo][idx] for mo in months], "total")
        if idx == 1:
            put("1元计价（补差价等）", [one_yuan[mo][idx] for mo in months], "aux")
            put("后台合计", [totals[mo][idx] + one_yuan[mo][idx] for mo in months], "raw")
            mom = [None] + [(totals[mo][1] / totals[months[i]][1] - 1) if totals[months[i]][1] else None
                            for i, mo in enumerate(months[1:])]
            ws.append(["主表合计环比"] + mom)
            for c in ws[ws.max_row][:ncol]:
                c.font = Font(name=FONT, size=9, color=C_MUTED)
                c.number_format = "+0.0%;-0.0%;0.0%"
        if last_cat >= first:  # 分类 × 月 色阶，一眼看出淡旺季和主力分类
            ws.conditional_formatting.add(
                f"B{first}:{openpyxl.utils.get_column_letter(1 + len(months))}{last_cat}",
                ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color=C_BAR))

    ops = read_ops(wb)
    if any(f"{year}-{mo:02d}" in ops for mo in months):
        ops_section(ws, year, months, ops, ncol)
    section("实际销售额（元）", 1, "#,##0", True)
    section("实际销量（件）", 0, "#,##0", False)
    ws.column_dimensions["A"].width = 22
    for i in range(len(months)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(2 + i)].width = 11
    ws.column_dimensions[openpyxl.utils.get_column_letter(2 + len(months))].width = 13
    ws.column_dimensions[openpyxl.utils.get_column_letter(3 + len(months))].width = 12
    ws.freeze_panes = "B3"


def ops_section(ws, year, months, ops, ncol):
    """经营指标：行 = 指标，列 = 月份 + 全年。全年列：可加的求和，比率按全年重算，余额取最后一个月。"""
    M = [ops.get(f"{year}-{mo:02d}") for mo in months]
    g = lambda k: [(r or {}).get(k) or 0 for r in M]
    S = lambda k: sum(g(k))
    div = lambda a, b: (a / b) if b else None
    rows = [
        ("订单与客单", None, None, None),
        ("订单数", g("订单数"), S("订单数"), "#,##0"),
        ("日均单", [div(r["订单数"], r["营业天数"]) if r else None for r in M], div(S("订单数"), S("营业天数")), "0.0"),
        ("客单价（元/单）", [div(r["订单实收"], r["订单数"]) if r else None for r in M], div(S("订单实收"), S("订单数")), "0.0"),
        ("件单价（元/件）", [div(r["主表销售额"], r["主表件数"]) if r else None for r in M], div(S("主表销售额"), S("主表件数")), "0.0"),
        ("新客 / 老客", None, None, None),
        ("老客单占比", [div(r["老客单数"], r["新客单数"] + r["老客单数"]) if r else None for r in M],
         div(S("老客单数"), S("新客单数") + S("老客单数")), "0.0%"),
        ("新客单占比", [div(r["新客单数"], r["新客单数"] + r["老客单数"]) if r else None for r in M],
         div(S("新客单数"), S("新客单数") + S("老客单数")), "0.0%"),
        ("储值", None, None, None),
        ("储值充值（实付）", g("储值充值"), S("储值充值"), "#,##0"),
        ("余额消费", g("余额消费"), S("余额消费"), "#,##0"),
        ("月末储值余额", g("期末储值余额"), next((v for v in reversed(g("期末储值余额")) if v), None), "#,##0"),
        ("营收与现金", None, None, None),
        ("订单实收（营业收入）", g("订单实收"), S("订单实收"), "#,##0"),
        ("　其中余额支付", g("余额支付"), S("余额支付"), "#,##0"),
        ("当月现金流入", [(r["订单实收"] - r["余额支付"] + r["储值充值"]) if r else None for r in M],
         S("订单实收") - S("余额支付") + S("储值充值"), "#,##0"),
        ("商品销售额（主表）", g("主表销售额"), S("主表销售额"), "#,##0"),
    ]
    ws.append([])
    ws.append(["经营指标（已剔除确认的异常）"])
    ws.cell(ws.max_row, 1).font = Font(name=FONT, bold=True, size=12, color=C_HEAD)
    ws.append(["指标"] + [f"{mo}月" for mo in months] + ["全年", ""])
    style_header(ws[ws.max_row][:ncol - 1])
    ws.row_dimensions[ws.max_row].height = 26
    for label, vals, yr, fmt in rows:
        if vals is None:  # 小标题
            ws.append([label])
            c = ws.cell(ws.max_row, 1)
            c.font = Font(name=FONT, bold=True, size=10, color=C_HEAD)
            for cc in ws[ws.max_row][:ncol - 1]:
                cc.fill = PatternFill("solid", fgColor=C_CAT)
            continue
        ws.append([label] + vals + [yr])
        row = ws[ws.max_row]
        for c in row[:ncol - 1]:
            c.font = Font(name=FONT, size=10)
            c.border = Border(bottom=THIN)
        for c in row[1:ncol - 1]:
            c.number_format = fmt
        row[ncol - 2].font = Font(name=FONT, bold=True, size=10)
    ws.append(["注：当月现金流入 = 订单实收 − 余额支付 + 储值充值；储值充值不算营收。月末余额还含充值赠送、后台充值、退款回余额，所以不等于「充值 − 消费」的累计。明细见「月度经营数据」。"])
    ws.cell(ws.max_row, 1).font = Font(name=FONT, italic=True, size=9, color=C_MUTED)


def restyle_existing_headers(wb):
    """已生成的月表 / 异常表：只调列宽，避免表头被挤成两行（不动数据）。"""
    for ws in wb.worksheets:
        mt = SHEET_RE.match(ws.title)
        if not mt:
            continue
        if mt.group(3) == "-异常与调整":
            ws.column_dimensions["F"].width = max(ws.column_dimensions["F"].width or 0, 12)
            ws.column_dimensions["G"].width = max(ws.column_dimensions["G"].width or 0, 14)
        else:
            ws.column_dimensions["E"].width = max(ws.column_dimensions["E"].width or 0, 12)
            ws.column_dimensions["G"].width = max(ws.column_dimensions["G"].width or 0, 19)


def rebuild_summary_only(path, rules):
    wb = openpyxl.load_workbook(path)
    year = int(re.match(r"^(\d{4})年", pathlib.Path(path).name).group(1))
    bak_dir = pathlib.Path(path).parent / "_backup"
    bak_dir.mkdir(exist_ok=True)
    shutil.copy2(path, bak_dir / f"{pathlib.Path(path).stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx")
    ops_rows = read_ops_rows(wb)
    if ops_rows:  # 按新格式重写，数字原样写回
        write_ops_sheet(wb, ops_rows)
    restyle_existing_headers(wb)
    write_year_summary(wb, year, rules.get("category_order") or [])
    if OPS_SHEET in wb.sheetnames:
        wb.move_sheet(OPS_SHEET, offset=1 - wb.sheetnames.index(OPS_SHEET))
    wb.active = 0
    wb.save(path)
    print(f"已重建 {path} 的「{SUMMARY_SHEET}」" + (f"和「{OPS_SHEET}」" if ops_rows else ""))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--summary-only", metavar="YEAR_XLSX", help="只重建该年度文件的「年度汇总」，不需要 --raw")
    ap.add_argument("--raw")
    ap.add_argument("--rules", default=str(pathlib.Path(__file__).resolve().parent.parent / "rules.yaml"))
    ap.add_argument("--dir", default=str(pathlib.Path.home() / "Desktop" / "五感专题分析"), help="年度文件所在目录（默认 ~/Desktop/五感专题分析）")
    ap.add_argument("--workbook", help="默认 <dir>/YYYY年五感销售数据统计.xlsx（一年一个文件，不存在就新建）")
    ap.add_argument("--sheet", help="默认 YYYY年M月")
    ap.add_argument("--dry-run", action="store_true", help="只校验+打印摘要，不写文件")
    ap.add_argument("--replace", action="store_true", help="同名 sheet 已存在时覆盖（仍会先备份）")
    a = ap.parse_args()

    rules = yaml.safe_load(pathlib.Path(a.rules).read_text(encoding="utf-8"))
    if a.summary_only:
        return rebuild_summary_only(a.summary_only, rules)
    if not a.raw:
        ap.error("需要 --raw（或用 --summary-only）")
    raw = load_raw(a.raw)
    y, m = raw["month"].split("-")
    sheet = a.sheet or f"{y}年{int(m)}月"

    items, ex, pending, anomaly_rows = apply_anomalies(raw, rules)
    if pending:
        print("ANOMALY_PENDING — 以下单笔 ≥ 阈值的消费还没确认，请问用户是否真实消费；"
              "确认后写进 rules.yaml 的 anomaly.excluded_orders / excluded_balance_pairs / kept：")
        for x in pending:
            print("  " + x)
        sys.exit(5)
    raw = dict(raw, list=items)
    rows, unknown = clean(raw, rules)
    if unknown:
        print("UNKNOWN_SPEC_OPTIONS — 以下规格选项未在 rules.yaml spec_rules 登记，请确认归 keep 还是 drop：")
        for (cat, o), n in unknown.most_common():
            print(f"  [{cat}] {o}  （{n} 行）")
        sys.exit(2)

    thr = float(rules.get("one_yuan_price_max", 1))
    one_rows = [r for r in rows if 0 < r["one_price"] <= thr]  # 标价 0 的是赠品，留在主表
    main_rows = [r for r in rows if not (0 < r["one_price"] <= thr)]
    cat_order = rules.get("category_order") or []

    # 对数：主表 + 1元 sheet = 后台明细合计（分）
    raw_qty = sum(int(float(r["real_num"] or 0)) for r in raw["list"])
    raw_c = sum(cents(r["real_total"]) for r in raw["list"])
    got_qty = sum(r["real_num"] for r in rows)
    got_c = sum(cents(r["real_total"]) for r in rows)
    if (raw_qty, raw_c) != (got_qty, got_c):
        print(f"RECON_FAIL 后台明细 {raw_qty} 件 / {raw_c/100:.2f} 元 vs 清洗后 {got_qty} 件 / {got_c/100:.2f} 元")
        sys.exit(3)
    hdr = {t["title"]: t["value"] for t in raw.get("header_totals", [])}
    hdr_real = hdr.get("实际销售额(元)")
    hdr_diff = round(float(hdr_real) - raw_c / 100 - ex["order_amt"], 2) if hdr_real is not None else None  # 页面顶部是剔除前口径

    main_lines = flatten(aggregate(main_rows, cat_order))
    one_lines = flatten(aggregate(one_rows, cat_order)) if one_rows else []

    src = pathlib.Path(a.workbook) if a.workbook else pathlib.Path(a.dir) / FILE_PATTERN.format(year=y)
    is_new = not src.exists()
    if is_new:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    else:
        wb = openpyxl.load_workbook(src)
    others = [openpyxl.load_workbook(f, read_only=True) for f in sorted(src.parent.glob("*五感销售数据统计*.xlsx"))
              if f.resolve() != src.resolve() and not f.name.startswith("~$")]
    prods, specs = history_names([wb] + others, (int(y), int(m)))

    # 异常与调整
    issues = []
    for r in rows:
        if r["orig_cat"] != r["cat"]:
            issues.append(["改分类", r["orig_cat"], r["orig_title"], "", f"→ {r['cat']}", r["real_num"], r["real_total"]])
        if r["orig_title"].strip() != r["title"]:
            issues.append(["改商品名", r["cat"], r["orig_title"], "", f"→ {r['title']}", r["real_num"], r["real_total"]])
        if r["orig_title"] != r["orig_title"].strip():
            issues.append(["去首尾空格", r["cat"], repr(r["orig_title"]), "", "", r["real_num"], r["real_total"]])
        if r["real_num"] > 0 and cents(r["real_total"]) == 0:
            issues.append(["0元成交(全额优惠/赠送)", r["cat"], r["title"], r["options"], f"标价合计 {r['sum_total']:.2f} 优惠 {r['discount']:.2f}", r["real_num"], 0])
    renamed_specs = {k for k in (rules.get("spec_rename") or {})}
    for r in rows:
        if r["spec"] and any(o in renamed_specs for o in r["options"].split("|")):
            issues.append(["改规格名", r["cat"], r["title"], r["options"], f"→ {r['spec']}", r["real_num"], r["real_total"]])
    for r in one_rows:
        issues.append(["1元计价→单独sheet", r["cat"], r["title"], r["options"], "", r["real_num"], r["real_total"]])
    seen_new = set()
    if prods:
        for r in main_rows:
            if r["title"] not in prods and r["title"] not in seen_new:
                seen_new.add(r["title"])
                issues.append(["新商品(此前月份未出现)", r["cat"], r["title"], "", "", "", ""])
            if r["spec"] and (r["title"], r["spec"]) not in specs and r["title"] in prods and (r["title"], r["spec"]) not in seen_new:
                seen_new.add((r["title"], r["spec"]))
                issues.append(["新规格(此前月份未出现)", r["cat"], r["title"], r["spec"], "", "", ""])
    for c in {r["cat"] for r in rows if r["cat"] not in cat_order}:
        issues.append(["新分类(未在category_order)", c, "", "", "", "", ""])
    issues = anomaly_rows + issues

    main_qty = sum(r["real_num"] for r in main_rows)
    main_c = sum(cents(r["real_total"]) for r in main_rows)
    print(f"{sheet}: 后台明细 {raw_qty} 件 / {raw_c/100:,.2f} 元（{len(raw['list'])} 行）")
    print(f"  主表 {main_qty} 件 / {main_c/100:,.2f} 元；1元计价 {raw_qty-main_qty} 件 / {(raw_c-main_c)/100:,.2f} 元")
    if ex["notes"]:
        print("  已剔除异常：" + "；".join(ex["notes"]) + (f"；成对充值 {ex['topup_amt']:,.0f}" if ex["topup_amt"] else ""))
    if hdr_diff:
        print(f"  注意：后台页面顶部「实际销售额」{hdr_real} 与明细合计差 {hdr_diff}（退款归属差异，主表以明细为准）")
    kinds = collections.Counter(i[0] for i in issues)
    print("  异常与调整：", dict(kinds) or "无")
    for i in issues:
        if i[0].startswith("新"):
            print("   ", i[0], i[1], i[2], i[3])

    if a.dry_run:
        return
    if sheet in wb.sheetnames and not a.replace:
        print(f"SHEET_EXISTS {sheet} 已存在，确认要覆盖再加 --replace")
        sys.exit(4)

    bak = None
    if not is_new:
        bak_dir = src.parent / "_backup"
        bak_dir.mkdir(exist_ok=True)
        bak = bak_dir / f"{src.stem}_{dt.datetime.now():%Y%m%d_%H%M%S}{src.suffix}"
        shutil.copy2(src, bak)

    for name in (sheet, f"{sheet}-1元计价", f"{sheet}-异常与调整"):
        if name in wb.sheetnames:
            del wb[name]
    note = (f"口径：支付时间 {raw['start']} ~ {raw['end']}；单价=后台标价均值；"
            f"不含 1 元计价商品（见「{sheet}-1元计价」），占比分母为本表合计；取数 {raw['fetched_at']}")
    write_sheet(wb.create_sheet(sheet), main_lines, note)
    if one_lines:
        write_sheet(wb.create_sheet(f"{sheet}-1元计价"), one_lines, "补差价/自定义金额类商品：后台单价 1 元，销量=金额，不代表杯数/件数")
    recon = ([("已剔除异常订单", f"{ex['order_n']} 笔 / {ex['order_amt']:,.2f} 元（见上方明细）")] if ex["order_n"] else []) + [
             ("后台明细合计（剔除后）", f"{raw_qty:,} 件 / {raw_c/100:,.2f} 元"),
             ("主表 + 1元计价", f"{got_qty:,} 件 / {got_c/100:,.2f} 元  ✓ 一致")]
    if hdr_real is not None:
        recon.append(("后台页面顶部实际销售额", f"{float(hdr_real):,.2f} 元" + (f"（与明细差 {hdr_diff}，以明细为准）" if hdr_diff else "  ✓ 一致")))
    write_issues(wb.create_sheet(f"{sheet}-异常与调整"), issues, recon)
    ops = ops_row(raw, ex, main_qty, main_c / 100)
    if ops:
        upsert_ops(wb, ops)
    else:
        print("  注意：raw 里没有营业报表/余额流水（旧版取数），本月不更新「月度经营数据」")
    sort_month_sheets(wb)
    write_year_summary(wb, int(y), cat_order)
    if OPS_SHEET in wb.sheetnames:  # 顺序：年度汇总 → 月度经营数据 → 各月
        wb.move_sheet(OPS_SHEET, offset=1 - wb.sheetnames.index(OPS_SHEET))
    wb.active = 0
    wb.save(src)
    print(f"已写入 {src} → sheet「{sheet}」；" + (f"备份 {bak}" if bak else "新建文件"))


if __name__ == "__main__":
    main()
