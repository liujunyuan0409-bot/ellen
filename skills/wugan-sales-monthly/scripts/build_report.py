"""五感月度销售表：raw json（fetch_month.js 输出）+ rules.yaml → 追加到工作簿的当月 sheet。

退出码：0 成功；2 有未登记的规格选项（需改 rules.yaml）；3 对数不平；4 sheet 已存在。
"""
import argparse
import collections
import datetime as dt
import json
import pathlib
import shutil
import sys

import openpyxl
import yaml
from openpyxl.styles import Font, PatternFill

HEADER = ["分类", "商品名称", "规格", "单价", "实际销量", "实际销售额(不含退款、优惠)", "实际销售额占比"]
BLANK_SPEC = "(空白)"


def cents(x):
    return round(float(x or 0) * 100)


def load_raw(path):
    d = json.loads(pathlib.Path(path).read_text(encoding="utf-8"))
    if isinstance(d, str):  # Playwright 落盘有时是 JSON 字符串套 JSON
        d = json.loads(d)
    if "error" in d:
        sys.exit(f"raw 文件是取数失败的结果：{d}")
    return d


def clean(raw, rules):
    spec_rules = rules.get("spec_rules") or {}
    unknown = collections.Counter()
    out = []
    for r in raw["list"]:
        title = (r["title"] or "").strip()
        cat = (rules.get("reclassify") or {}).get(title, r["classify_title"])
        orig_cat = r["classify_title"]
        title = (rules.get("product_rename") or {}).get(title, title)
        spec = None
        sr = spec_rules.get(orig_cat) or spec_rules.get(cat)
        if sr:
            opts = [o.strip() for o in r["key_list"] if o and o.strip()]
            for o in opts:
                if o not in sr.get("keep", []) and o not in sr.get("drop", []):
                    unknown[(orig_cat, o)] += 1
            kept = [o for o in opts if o in sr.get("keep", [])]
            if kept:
                spec = (rules.get("spec_rename") or {}).get(kept[0], kept[0])
        out.append({
            "cat": cat, "orig_cat": orig_cat, "orig_title": r["title"], "title": title, "spec": spec,
            "one_price": float(r["one_price"] or 0), "real_num": int(float(r["real_num"] or 0)),
            "real_total": float(r["real_total"] or 0), "discount": float(r["discount_amount"] or 0),
            "sum_total": float(r["sum_total"] or 0), "options": "|".join(r["key_list"]),
        })
    return out, unknown


def aggregate(rows, cat_order):
    """三层：分类 → 商品 → 规格。单价 = 明细行 one_price 的算术平均。"""
    def node():
        return {"prices": [], "qty": 0, "amt_c": 0, "kids": collections.OrderedDict()}

    root = collections.OrderedDict()
    for r in rows:
        c = root.setdefault(r["cat"], node())
        p = c["kids"].setdefault(r["title"], node())
        targets = [c, p]
        if r["spec"] is not None or any(x["spec"] for x in rows if x["cat"] == r["cat"] and x["title"] == r["title"]):
            targets.append(p["kids"].setdefault(r["spec"] or BLANK_SPEC, node()))
        for n in targets:
            n["prices"].append(r["one_price"])
            n["qty"] += r["real_num"]
            n["amt_c"] += cents(r["real_total"])

    def order_key(name):
        return (cat_order.index(name) if name in cat_order else len(cat_order), name)

    def sort_kids(n):
        items = sorted(n["kids"].items(), key=lambda kv: (kv[0] == BLANK_SPEC, -kv[1]["amt_c"], kv[0]))
        n["kids"] = collections.OrderedDict(items)
        for _, k in items:
            sort_kids(k)

    ordered = collections.OrderedDict(sorted(root.items(), key=lambda kv: order_key(kv[0])))
    for n in ordered.values():
        sort_kids(n)
    return ordered


def flatten(tree):
    total_c = sum(n["amt_c"] for n in tree.values())
    lines = []
    for cname, c in tree.items():
        lines.append((0, [cname, None, None], c))
        for pname, p in c["kids"].items():
            lines.append((1, [None, pname, None], p))
            for sname, s in p["kids"].items():
                lines.append((2, [None, None, sname], s))
    out = []
    for level, names, n in lines:
        price = sum(n["prices"]) / len(n["prices"]) if n["prices"] else None
        share = n["amt_c"] / total_c if total_c else 0
        out.append((level, names + [price, n["qty"], n["amt_c"] / 100, share]))
    all_prices = [x for n in tree.values() for x in n["prices"]]
    out.append((-1, ["总计", None, None, sum(all_prices) / len(all_prices) if all_prices else None,
                     sum(n["qty"] for n in tree.values()), total_c / 100, 1 if total_c else 0]))
    return out


def history_names(wb):
    """已有 sheet 里出现过的 (商品名) 和 (商品名, 规格)，用来识别新品。"""
    prods, specs = set(), set()
    for ws in wb.worksheets:
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        rows = [r for r in rows if any(v is not None for v in r)]
        if not rows:
            continue
        head = [str(v).strip() if v is not None else "" for v in rows[0]]
        if "商品名称" not in head or "规格" not in head:
            continue
        bi, ci = head.index("商品名称"), head.index("规格")
        cur = None
        for r in rows[1:]:
            if r[bi]:
                cur = str(r[bi]).strip()
                prods.add(cur)
            elif r[ci] and cur:
                specs.add((cur, str(r[ci]).strip()))
    return prods, specs


def write_sheet(ws, lines, note=None):
    ws.append(HEADER)
    for c in ws[1]:
        c.font = Font(bold=True)
    cat_fill = PatternFill("solid", fgColor="FFF2CC")
    for level, vals in lines:
        ws.append(vals)
        row = ws[ws.max_row]
        if level in (0, -1):
            for c in row:
                c.font = Font(bold=True)
                if level == 0:
                    c.fill = cat_fill
        row[3].number_format = "0.00"
        row[5].number_format = "#,##0.00"
        row[6].number_format = "0.00%"
    if note:
        ws.append([])
        ws.append([note])
    for col, w in zip("ABCDEFG", [20, 32, 14, 10, 10, 24, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", required=True)
    ap.add_argument("--rules", default=str(pathlib.Path(__file__).resolve().parent.parent / "rules.yaml"))
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--sheet", help="默认 YYYY年M月")
    ap.add_argument("--dry-run", action="store_true", help="只校验+打印摘要，不写文件")
    ap.add_argument("--replace", action="store_true", help="同名 sheet 已存在时覆盖（仍会先备份）")
    a = ap.parse_args()

    raw = load_raw(a.raw)
    rules = yaml.safe_load(pathlib.Path(a.rules).read_text(encoding="utf-8"))
    y, m = raw["month"].split("-")
    sheet = a.sheet or f"{y}年{int(m)}月"

    rows, unknown = clean(raw, rules)
    if unknown:
        print("UNKNOWN_SPEC_OPTIONS — 以下规格选项未在 rules.yaml spec_rules 登记，请确认归 keep 还是 drop：")
        for (cat, o), n in unknown.most_common():
            print(f"  [{cat}] {o}  （{n} 行）")
        sys.exit(2)

    thr = float(rules.get("one_yuan_price_max", 1))
    one_rows = [r for r in rows if 0 < r["one_price"] <= thr]  # 标价 0 的是赠品，留在主表
    main_rows = [r for r in rows if not (0 < r["one_price"] <= thr)]
    cat_order = rules.get("category_order") or []

    # 对数：主表 + 1元 sheet = 后台明细合计（分）
    raw_qty = sum(int(float(r["real_num"] or 0)) for r in raw["list"])
    raw_c = sum(cents(r["real_total"]) for r in raw["list"])
    got_qty = sum(r["real_num"] for r in rows)
    got_c = sum(cents(r["real_total"]) for r in rows)
    if (raw_qty, raw_c) != (got_qty, got_c):
        print(f"RECON_FAIL 后台明细 {raw_qty} 件 / {raw_c/100:.2f} 元 vs 清洗后 {got_qty} 件 / {got_c/100:.2f} 元")
        sys.exit(3)
    hdr = {t["title"]: t["value"] for t in raw.get("header_totals", [])}
    hdr_real = hdr.get("实际销售额(元)")
    hdr_diff = round(float(hdr_real) - raw_c / 100, 2) if hdr_real is not None else None

    main_lines = flatten(aggregate(main_rows, cat_order))
    one_lines = flatten(aggregate(one_rows, cat_order)) if one_rows else []

    wb = openpyxl.load_workbook(a.workbook)
    prods, specs = history_names(wb)

    # 异常与调整
    issues = []
    for r in rows:
        if r["orig_cat"] != r["cat"]:
            issues.append(["改分类", r["orig_cat"], r["orig_title"], "", f"→ {r['cat']}", r["real_num"], r["real_total"]])
        if r["orig_title"].strip() != r["title"]:
            issues.append(["改商品名", r["cat"], r["orig_title"], "", f"→ {r['title']}", r["real_num"], r["real_total"]])
        if r["orig_title"] != r["orig_title"].strip():
            issues.append(["去首尾空格", r["cat"], repr(r["orig_title"]), "", "", r["real_num"], r["real_total"]])
        if r["real_num"] > 0 and cents(r["real_total"]) == 0:
            issues.append(["0元成交(全额优惠/赠送)", r["cat"], r["title"], r["options"], f"标价合计 {r['sum_total']:.2f} 优惠 {r['discount']:.2f}", r["real_num"], 0])
    renamed_specs = {k for k in (rules.get("spec_rename") or {})}
    for r in rows:
        if r["spec"] and any(o in renamed_specs for o in r["options"].split("|")):
            issues.append(["改规格名", r["cat"], r["title"], r["options"], f"→ {r['spec']}", r["real_num"], r["real_total"]])
    for r in one_rows:
        issues.append(["1元计价→单独sheet", r["cat"], r["title"], r["options"], "", r["real_num"], r["real_total"]])
    seen_new = set()
    if prods:
        for r in main_rows:
            if r["title"] not in prods and r["title"] not in seen_new:
                seen_new.add(r["title"])
                issues.append(["新商品(历史sheet未出现)", r["cat"], r["title"], "", "", "", ""])
            if r["spec"] and (r["title"], r["spec"]) not in specs and r["title"] in prods and (r["title"], r["spec"]) not in seen_new:
                seen_new.add((r["title"], r["spec"]))
                issues.append(["新规格(历史sheet未出现)", r["cat"], r["title"], r["spec"], "", "", ""])
    for c in {r["cat"] for r in rows if r["cat"] not in cat_order}:
        issues.append(["新分类(未在category_order)", c, "", "", "", "", ""])

    main_qty = sum(r["real_num"] for r in main_rows)
    main_c = sum(cents(r["real_total"]) for r in main_rows)
    print(f"{sheet}: 后台明细 {raw_qty} 件 / {raw_c/100:,.2f} 元（{len(raw['list'])} 行）")
    print(f"  主表 {main_qty} 件 / {main_c/100:,.2f} 元；1元计价 {raw_qty-main_qty} 件 / {(raw_c-main_c)/100:,.2f} 元")
    if hdr_diff:
        print(f"  注意：后台页面顶部「实际销售额」{hdr_real} 与明细合计差 {hdr_diff}（退款归属差异，主表以明细为准）")
    kinds = collections.Counter(i[0] for i in issues)
    print("  异常与调整：", dict(kinds) or "无")
    for i in issues:
        if i[0].startswith("新"):
            print("   ", i[0], i[1], i[2], i[3])

    if a.dry_run:
        return
    if sheet in wb.sheetnames and not a.replace:
        print(f"SHEET_EXISTS {sheet} 已存在，确认要覆盖再加 --replace")
        sys.exit(4)

    src = pathlib.Path(a.workbook)
    bak_dir = src.parent / "_backup"
    bak_dir.mkdir(exist_ok=True)
    bak = bak_dir / f"{src.stem}_{dt.datetime.now():%Y%m%d_%H%M%S}{src.suffix}"
    shutil.copy2(src, bak)

    for name in (sheet, f"{sheet}-1元计价", f"{sheet}-异常与调整"):
        if name in wb.sheetnames:
            del wb[name]
    note = (f"口径：支付时间 {raw['start']} ~ {raw['end']}；单价=后台标价均值；"
            f"不含 1 元计价商品（见「{sheet}-1元计价」），占比分母为本表合计；取数 {raw['fetched_at']}")
    write_sheet(wb.create_sheet(sheet), main_lines, note)
    if one_lines:
        write_sheet(wb.create_sheet(f"{sheet}-1元计价"), one_lines, "补差价/自定义金额类商品：后台单价 1 元，销量=金额，不代表杯数/件数")
    ws = wb.create_sheet(f"{sheet}-异常与调整")
    ws.append(["类型", "分类", "商品", "规格/选项", "说明", "实际销量", "实际销售额"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i in issues:
        ws.append(i)
    ws.append([])
    ws.append(["对数", "后台明细合计", "", "", f"{raw_qty} 件 / {raw_c/100:.2f} 元", "", ""])
    ws.append(["对数", "主表 + 1元计价", "", "", f"{got_qty} 件 / {got_c/100:.2f} 元", "", ""])
    if hdr_real is not None:
        ws.append(["对数", "后台页面顶部实际销售额", "", "", f"{hdr_real}（差 {hdr_diff}）", "", ""])
    for col, w in zip("ABCDEFG", [26, 20, 32, 28, 36, 10, 12]):
        ws.column_dimensions[col].width = w
    wb.save(src)
    print(f"已写入 {src} → sheet「{sheet}」；备份 {bak}")


if __name__ == "__main__":
    main()
